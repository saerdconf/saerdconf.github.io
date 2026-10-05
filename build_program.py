#!/usr/bin/env python3
"""Build program.qmd and speakers.qmd for the SAERD site from the program workbook.

Usage:  python build_program.py SAERD26_Program_v6.xlsx
Speaker photos: drop images/speakers/<slug>.jpg|jpeg|png|webp (slug = name without
titles, lower-case, hyphens; shown in the printed list below). Bios: fill the
'Bio' column on the workbook's 'Speakers' sheet and re-run.
"""
import sys, re, html, datetime as dt
from pathlib import Path
from openpyxl import load_workbook

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "SAERD26_Program_v6.xlsx")
ROOT = Path(__file__).parent
esc = html.escape

wb = load_workbook(SRC, data_only=True)

def table(ws, first_header):
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        if r and r[0] == first_header:
            hdr = [str(h).strip() if h is not None else "" for h in r]
            out = []
            for r2 in rows[i + 1:]:
                if all(v is None for v in r2):
                    continue
                out.append({h: v for h, v in zip(hdr, r2) if h})
            return out
    raise SystemExit(f"header '{first_header}' not found in {ws.title}")

def to_time(v):
    if isinstance(v, dt.datetime): return v.time()
    if isinstance(v, dt.time): return v
    m = round(float(v) * 24 * 60)
    return dt.time(m // 60, m % 60)

def to_date(v):
    if isinstance(v, dt.datetime): return v.date()
    if isinstance(v, dt.date): return v
    return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))

def hhmm(t): return t.strftime("%H:%M")

TITLES = re.compile(r"^(H\.E\.|Dr\.|Prof\.|Asso\.|Asst\.|Assoc\.|Mr\.|Miss\.|Ms\.|Mrs\.)\s*", re.I)
def bare(name):
    prev = None
    while prev != name:
        prev, name = name, TITLES.sub("", name.strip())
    return name
def slug(name):
    return re.sub(r"[^a-z0-9]+", "-", bare(name).lower()).strip("-")
def initials(name):
    parts = [p for p in bare(name).split() if p[0].isalpha()]
    return (parts[0][0] + (parts[-1][0] if len(parts) > 1 else "")).upper()

settings = {r[0]: r[1] for r in wb["Settings"].iter_rows(values_only=True) if r[0] and r[1] is not None}
CONF = settings.get("Conference name", "SAERD 2026")
VENUE = settings.get("Venue and location", "")
HASHTAG = settings.get("Hashtag", "")

speakers = [s for s in table(wb["Speakers"], "No") if s.get("Name")]
sessions = {s["Code"]: s for s in table(wb["Sessions"], "Code") if s.get("Code")}
papers = [p for p in table(wb["Data"], "No") if p.get("Session")]
schedule = [s for s in table(wb["Schedule"], "Day") if s.get("Item")]

def real(sp): return not sp["Name"].lower().startswith("speaker to be announced")
by_name = {sp["Name"]: sp for sp in speakers if real(sp)}

def person(name, with_aff=True):
    sp = by_name.get(name)
    if not sp:
        return f'<span class="person">{esc(name)}</span>'
    aff = f' <span class="aff">{esc(sp.get("Position / affiliation") or "")}</span>' if with_aff else ""
    return f'<span class="person"><a href="speakers.html#{slug(name)}">{esc(name)}</a>{aff}</span>'

# --- plenary details (structured; names matching the Speakers sheet are linked) ---
PLENARY = {
  "Edited Book Project": ("Asian Developing Region", [("", ["Prof. Sovannroeun Samreth", "Dr. Sattwick Dey Biswas"])]),
  "Opening and welcome remarks": ("", [("", ["Dr. Ratha Chiv", "H.E. Dr. Kao Thach"])]),
  "Panel Discussion A": ("ARDB: Agriculture and Rural Financing", [("Speakers", ["To be announced"]), ("Moderator", ["To be announced"])]),
  "Introduction to the Young Scholar Initiative": ("", [("", ["Sokhimmarya Chea (YSI)", "Sinoun Him (YSI)", "Jenny Symaly (YSI)"])]),
  "Panel Discussion B": ("Economic Growth and Sectoral Transformation in Developing Economies",
      [("Speakers", ["Prof. Matthew McCartney", "Prof. Micheal Hill", "Asso. Prof. Thirunaukarasu Subramaniam", "Asso. Prof. Penghuy Ngov"]),
       ("Moderator", ["Prof. Sovannroeun Samreth"])]),
  "Special session": ("Master Thesis Development Award", []),
  "Closing remarks": ("", [("", ["The Conference Organizing Committee"])]),
}

def slot(row):
    t0, t1 = hhmm(to_time(row["Start"])), hhmm(to_time(row["End"]))
    item, kind = row["Item"], (row.get("Type") or "").lower()
    body = f'<div class="slot-title">{esc(item)}</div>'
    if item in PLENARY:
        sub, groups = PLENARY[item]
        if sub: body += f'<div class="slot-sub">{esc(sub)}</div>'
        for label, names in groups:
            if label: body += f'<div class="slot-role">{esc(label)}</div>'
            body += '<div class="slot-people">' + "".join(person(n) for n in names) + "</div>"
    elif kind == "parallel":
        body += '<div class="slot-sub">Five sessions run at the same time</div>'
    elif row.get("Details"):
        body += f'<div class="slot-sub">{esc(" ".join(str(row["Details"]).split()))}</div>'
    return f'<div class="slot slot-{kind or "other"}"><div class="slot-time">{t0} – {t1}</div><div class="slot-body">{body}</div></div>'

def paper_li(p):
    printed = " ".join(str(p.get("Authors (as printed)") or "").split())
    lead = f'{p["First name"]} {p["Last name"]}'.strip()
    if printed.startswith(lead):
        printed = f"<b>{esc(lead)}</b>{esc(printed[len(lead):])}"
    else:
        printed = esc(printed)
    return f'<li><div class="paper-title">{esc(str(p["Paper title"]).strip())}</div><div class="paper-authors">{printed}</div></li>'

def card(code):
    s = sessions[code]
    ps = sorted([p for p in papers if p["Session"] == code], key=lambda p: p["Order"])
    thesis = "thesis" if s["Session title"].startswith("Thesis Development") else ""
    badge = '<span class="badge">Master’s Thesis Workshop</span>' if thesis else ""
    return (f'<section class="session-card{" " + thesis if thesis else ""}"><header><span class="code">{code}</span>'
            f'<span class="room">{esc(s["Room"])}</span>{badge}</header>'
            f'<div class="session-title">{esc(s["Session title"])}</div>'
            f'<div class="chair"><b>Chair:</b> {esc(s["Chair"])}</div>'
            f'<ol class="papers">{"".join(paper_li(p) for p in ps)}</ol></section>')

def parallel(row):
    block = row["Item"].split()[-1]
    codes = sorted([c for c, s in sessions.items() if s["Block"] == block], key=lambda c: sessions[c]["Room #"])
    return slot(row).replace(esc(row["Item"]), esc(row["Item"]), 1) + \
           '<div class="parallel-grid">' + "".join(card(c) for c in codes) + "</div>"

# --- program.qmd ---
days = sorted({s["Day"] for s in schedule})
dates = {}
for r in schedule: dates[r["Day"]] = to_date(r["Date"])
out = ["---", 'title: "Conference Program"', "css: assets/program.css", "toc: false", "page-layout: full", "---", "",
       f'**{esc(CONF)}**  ', f'**{dates[days[0]].day}–{dates[days[-1]].day} {dates[days[-1]]:%B %Y}** · {esc(VENUE)} · {esc(HASHTAG)}', "",
       "::: {.panel-tabset}", ""]
for d in days:
    out += [f"### Day {d} · {dates[d]:%a} {dates[d].day} {dates[d]:%b}", "", "```{=html}", f'<div class="program-day">']
    for r in [x for x in schedule if x["Day"] == d]:
        out.append(parallel(r) if (r.get("Type") or "") == "Parallel" else slot(r))
    out += ["</div>", "```", ""]
out += [":::", "", f"[Meet the speakers →](speakers.qmd)", ""]
(ROOT / "program.qmd").write_text("\n".join(out), encoding="utf-8")

# --- speakers.qmd ---
def photo(name):
    for ext in ("jpg", "jpeg", "png", "webp"):
        if (ROOT / "images" / "speakers" / f"{slug(name)}.{ext}").exists():
            return f'<img src="images/speakers/{slug(name)}.{ext}" alt="{esc(name)}" loading="lazy">'
    return f'<div class="initials" aria-hidden="true">{esc(initials(name))}</div>'

cards = []
for sp in speakers:
    if not real(sp): continue
    n = sp["Name"]
    bio = (sp.get("Bio") or "").strip()
    bio_html = "".join(f"<p>{esc(p.strip())}</p>" for p in bio.split("\n") if p.strip()) if bio \
               else '<p class="bio-tba">Biography to be announced.</p>'
    cards.append(f'<article class="speaker-card" id="{slug(n)}"><div class="speaker-photo">{photo(n)}</div>'
                 f'<div class="speaker-info"><h3>{esc(n)}</h3><div class="speaker-aff">{esc(sp.get("Position / affiliation") or "")}</div>'
                 f'<div class="speaker-role">{esc(sp.get("Role at the conference") or "")}</div>{bio_html}</div></article>')
sq = ["---", 'title: "Speakers"', "css: assets/program.css", "toc: false", "page-layout: full", "---", "",
      "Opening remarks, panel discussants, and plenary speakers at SAERD 2026. Panel A speakers will be announced.", "",
      "```{=html}", '<div class="speaker-grid">', *cards, "</div>", "```", "",
      "[View the conference program →](program.qmd)", ""]
(ROOT / "speakers.qmd").write_text("\n".join(sq), encoding="utf-8")

print(f"program.qmd: {len(papers)} papers in {len(sessions)} sessions; speakers.qmd: {len(cards)} speakers")
for sp in speakers:
    if real(sp): print("  photo slug:", slug(sp["Name"]))
